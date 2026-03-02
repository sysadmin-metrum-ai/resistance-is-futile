#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

SLIDES = [
    ("slide-01-title.png", "AI Ops Drone Swarm - Title"),
    ("slide-02-problem.png", "The Problem"),
    ("slide-03-solution.png", "The Solution"),
    ("slide-04-architecture.png", "System Architecture - AMD MI350P"),
    ("slide-05-value.png", "Value Proposition"),
    ("slide-06-status.png", "Milestone v1.0 Status"),
    ("slide-07-flow.png", "Runtime Flow"),
    ("slide-08-demo.png", "DTW 2026 Primary Demo"),
    ("slide-09-security.png", "Use Case: Physical Security"),
    ("slide-10-cable.png", "Use Case: Cable Inspection"),
    ("slide-11-maintenance.png", "Use Case: Post-Maintenance"),
    ("slide-12-timeline.png", "DTW 2026 Timeline"),
    ("slide-13-milestones.png", "Weekly Milestones"),
    ("slide-14-success.png", "Success Metrics & CTA"),
]

prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)

for i, (img_file, title) in enumerate(SLIDES, 1):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    img_path = f"docs/slide-images/{img_file}"

    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=Inches(16))

    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.2), Inches(15), Inches(0.8)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"Slide {i}: {title}"
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(150, 150, 150)

prs.save("docs/DTW-2026-Demo-Slide-Deck.pptx")
print("Created: docs/DTW-2026-Demo-Slide-Deck.pptx")
